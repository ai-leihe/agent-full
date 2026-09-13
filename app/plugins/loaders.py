"""loader 槽位插件：把上传的原始文件解码并抽取为纯文本。

同一槽位提供多个实现（text / pdf / docx），上游只需声明需要 ``text`` 产物，
实际用哪个 loader 由 YAML 决定。
"""

from __future__ import annotations

import csv
import html
import io
import json
import re
from pathlib import Path

from ..core.context import PipelineContext, RAW_FILE, TEXT
from ..core.skill import Skill, skill

#: 按优先级尝试的编码，覆盖中英文常见场景
_ENCODING_CANDIDATES = ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1")

_SCRIPT_RE = re.compile(r"<(script|style)\b.*?</\1>", re.IGNORECASE | re.DOTALL)
_TAG_RE = re.compile(r"<[^>]+>")
_BLANK_RE = re.compile(r"[ \t]+\n")

#: 文件签名（魔数）：后缀可以骗人，字节不会
_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"


def sniff_format(raw: bytes) -> str:
    """按文件签名判断真实格式，返回 ``pdf`` / ``zip`` / ``text``。"""
    if raw[:1024].lstrip().startswith(_PDF_MAGIC):
        return "pdf"
    if raw.startswith(_ZIP_MAGIC):
        return "zip"
    return "text"


def is_binary(raw: bytes) -> bool:
    """粗判二进制：NUL 字节或过高的不可打印字符占比。

    用它拦住「用 latin-1 硬解码二进制 → 得到乱码文本」这条路径。
    """
    sample = raw[:4096]
    if not sample:
        return False
    if b"\x00" in sample:
        return True
    odd = sum(1 for byte in sample if byte < 0x09 or 0x0E <= byte < 0x20)
    return odd / len(sample) > 0.10


def _looks_like_other_document(suffix: str, raw: bytes) -> bool:
    """后缀声明是文档格式，但内容既不匹配签名又像二进制 → 大概率是坏文件或改了扩展名。"""
    if suffix not in {".pdf", ".docx", ".doc", ".xlsx", ".pptx"}:
        return False
    return is_binary(raw)


def _clean_extracted(text: str) -> str:
    """压掉抽取结果里常见的 NUL、零宽字符与连续空行。"""
    text = text.replace("\x00", "").replace("\ufeff", "")
    text = _BLANK_RE.sub("\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf_text(raw: bytes) -> tuple[str, int]:
    """用 pypdf 逐页抽取 PDF 文本，返回 (文本, 页数)。"""
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:  # pragma: no cover
        raise RuntimeError("解析 PDF 需要额外依赖，请执行：pip install pypdf") from None

    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and reader.decrypt("") == 0:  # 0 = 空密码也解不开
            raise ValueError("PDF 已加密，无法抽取文本；请先解除密码保护再上传")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - 底层异常统一转成可读提示
        raise ValueError(
            f"PDF 解析失败：文件可能已损坏或被截断（{type(exc).__name__}）"
        ) from exc

    text = _clean_extracted("\n\n".join(page for page in pages if page))
    if not text:
        raise ValueError("PDF 未抽取到文本，可能是扫描件（图片型 PDF），需要 OCR 插件")
    return text, len(pages)


def extract_docx_text(raw: bytes) -> str:
    """用 python-docx 抽取段落与表格文本。"""
    try:
        import docx  # type: ignore
    except ImportError:  # pragma: no cover
        raise RuntimeError("解析 .docx 需要额外依赖，请执行：pip install python-docx") from None

    document = docx.Document(io.BytesIO(raw))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = _clean_extracted("\n".join(parts))
    if not text:
        raise ValueError("docx 未抽取到文本")
    return text


def decode_bytes(raw: bytes) -> tuple[str, str]:
    """把字节流解码为文本，返回 (文本, 使用的编码)。"""
    for encoding in _ENCODING_CANDIDATES:
        try:
            return raw.decode(encoding), encoding
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(replace)"


def _html_to_text(source: str) -> str:
    source = _SCRIPT_RE.sub(" ", source)
    source = re.sub(r"<br\s*/?>", "\n", source, flags=re.IGNORECASE)
    source = re.sub(r"</(p|div|li|tr|h[1-6])>", "\n", source, flags=re.IGNORECASE)
    source = _TAG_RE.sub(" ", source)
    return html.unescape(source)


def _csv_to_text(source: str) -> str:
    rows = list(csv.reader(io.StringIO(source)))
    if not rows:
        return source
    header, *body = rows
    lines = [" | ".join(header)]
    lines += [" | ".join(cell.strip() for cell in row) for row in body]
    return "\n".join(lines)


@skill
class TextLoader(Skill):
    """通用文本加载器：txt / md / csv / json / log / xml / html 等。"""

    name = "text_loader"
    slot = "loader"
    description = "把文件解析为规范化纯文本：文本类直接解码，PDF / Word 按文件签名自动委派抽取。"
    consumes = (RAW_FILE,)
    produces = (TEXT,)
    param_schema = {
        "encoding": {
            "type": "str",
            "default": "",
            "label": "文件编码",
            "help": "留空则自动探测",
            "choices": ["", "utf-8", "gb18030", "big5", "latin-1"],
        },
    }

    def configure(self, options: dict) -> None:
        self.encoding = options.get("encoding") or None

    def run(self, ctx: PipelineContext) -> None:
        raw: bytes = ctx.require(RAW_FILE)
        suffix = Path(ctx.filename).suffix.lower()

        # 二进制文档绝不能按文本解码（会得到 %PDF-1.7 / /Filter/FlateDecode 之类的乱码），
        # 这里按文件签名把它们交给真正的解析器，后缀不可靠时也能兜住。
        kind = sniff_format(raw)
        if kind == "pdf":
            text, pages = extract_pdf_text(raw)
            ctx.put(
                TEXT, text, producer=self.name, format="pdf",
                pages=pages, suffix=suffix, source_bytes=len(raw),
            )
            return
        if kind == "zip":
            if suffix in {".docx", ".docm"} or b"word/document.xml" in raw[:8192]:
                text = extract_docx_text(raw)
                ctx.put(
                    TEXT, text, producer=self.name, format="docx",
                    suffix=suffix, source_bytes=len(raw),
                )
                return
            raise ValueError(
                f"'{suffix or '未知格式'}' 属于压缩包类文档（xlsx / pptx / 压缩包），"
                "text_loader 无法解析；请改用对应的 loader 插件，或先另存为文本型文件"
            )
        if _looks_like_other_document(suffix, raw):
            raise ValueError(
                f"文件后缀是 {suffix}，但内容不是有效的 {suffix.lstrip('.').upper()}："
                "请确认文件没有被改过扩展名，或换成对应的 loader 插件"
            )

        if self.encoding:
            text = raw.decode(self.encoding, errors="replace")
            encoding = self.encoding
        else:
            text, encoding = decode_bytes(raw)

        if suffix in {".html", ".htm", ".xhtml"}:
            text = _html_to_text(text)
        elif suffix == ".csv" or ctx.content_type == "text/csv":
            text = _csv_to_text(text)
        elif suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError:
                pass

        text = _BLANK_RE.sub("\n", text).strip()
        if not text:
            raise ValueError("文件解析后为空，请确认文件内容与编码")

        ctx.put(TEXT, text, producer=self.name, encoding=encoding, suffix=suffix, source_bytes=len(raw))


@skill
class PdfLoader(Skill):
    """PDF 加载器（可选插件，依赖 pypdf）。"""

    name = "pdf_loader"
    slot = "loader"
    optional = True
    description = "从 PDF 中逐页抽取文本，需额外安装 pypdf。"
    consumes = (RAW_FILE,)
    produces = (TEXT,)

    def run(self, ctx: PipelineContext) -> None:
        text, pages = extract_pdf_text(ctx.require(RAW_FILE))
        ctx.put(TEXT, text, producer=self.name, pages=pages, scanned=False)


@skill
class DocxLoader(Skill):
    """Word 加载器（可选插件，依赖 python-docx）。"""

    name = "docx_loader"
    slot = "loader"
    optional = True
    description = "抽取 .docx 的段落与表格文本，需额外安装 python-docx。"
    consumes = (RAW_FILE,)
    produces = (TEXT,)

    def run(self, ctx: PipelineContext) -> None:
        text = extract_docx_text(ctx.require(RAW_FILE))
        ctx.put(TEXT, text, producer=self.name, format="docx")
